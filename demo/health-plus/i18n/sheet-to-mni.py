"""Turn the filled translation sheet into i18n/mni.js.

Usage: python3 sheet-to-mni.py "Health Plus copy - translation sheet.xlsx"
Reads the "Website copy" tab: column D is the key, column F the Romanised Meiteilon text.
Empty cells are skipped, so those lines stay in English.
"""
import json
import os
import sys

from openpyxl import load_workbook

ws = load_workbook(sys.argv[1], read_only=True)['Website copy']
out = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    key, mni = row[3], row[5]
    if key and mni and str(mni).strip():
        out[key] = str(mni).strip()

here = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(here, 'mni.js'), 'w', encoding='utf-8') as f:
    f.write('/* Romanised Meiteilon copy, generated from the translation sheet by sheet-to-mni.py.\n'
            '   Any key left out falls back to English. */\n')
    f.write('window.HP_I18N_MNI = ' + json.dumps(out, ensure_ascii=False, indent=1) + ';\n')
print(f'{len(out)} lines written to mni.js')
